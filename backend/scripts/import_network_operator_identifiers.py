from pathlib import Path

from openpyxl import load_workbook

from app.database.session import SessionLocal
from app.models.network_operator import NetworkOperator
from app.models.network_operator_identifier import NetworkOperatorIdentifier


DATA_DIR = (
    Path(__file__).resolve().parent.parent
    / "data"
    / "network_operators"
)

FILES = (
    {
        "pattern": "STROM_EC_ZPN_Praefixe_2026_mit_Bundesland*.xlsx",
        "sheet": "NB_EC_Import",
        "energy_type": "Strom",
    },
    {
        "pattern": "GAS_EC_ZPN_Praefixe_2026_mit_Bundesland*.xlsx",
        "sheet": "GAS_NB_EC_Import",
        "energy_type": "Gas",
    },
)

# Existing local names that intentionally differ from the source file.
OPERATOR_ALIASES = {
    "Wiener Netze GmbH": "Wiener Netze",
}


def clean(value) -> str:
    if value is None:
        return ""

    return str(value).strip()


def find_file(pattern: str) -> Path:
    matches = sorted(DATA_DIR.glob(pattern))

    if not matches:
        raise FileNotFoundError(
            f"Keine Datei gefunden: {DATA_DIR / pattern}"
        )

    if len(matches) > 1:
        raise RuntimeError(
            f"Mehrere Dateien gefunden für {pattern}: "
            + ", ".join(path.name for path in matches)
        )

    return matches[0]


def find_operator(
    db,
    source_name: str,
) -> NetworkOperator | None:
    local_name = OPERATOR_ALIASES.get(
        source_name,
        source_name,
    )

    return (
        db.query(NetworkOperator)
        .filter(NetworkOperator.name.ilike(local_name))
        .first()
    )


def get_or_create_operator(
    db,
    source_name: str,
) -> tuple[NetworkOperator, bool]:
    operator = find_operator(
        db,
        source_name,
    )

    if operator:
        if not operator.active:
            operator.active = True

        return operator, False

    operator = NetworkOperator(
        name=source_name,
        code=None,
        active=True,
    )

    db.add(operator)
    db.flush()

    return operator, True


def import_file(
    db,
    *,
    path: Path,
    sheet_name: str,
    energy_type: str,
) -> dict:
    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    if sheet_name not in workbook.sheetnames:
        raise RuntimeError(
            f"Sheet '{sheet_name}' fehlt in {path.name}"
        )

    sheet = workbook[sheet_name]

    created_operators = 0
    created_identifiers = 0
    updated_identifiers = 0
    processed = 0

    for row_number, row in enumerate(
        sheet.iter_rows(
            min_row=2,
            values_only=True,
        ),
        start=2,
    ):
        source_name = clean(row[0])
        zpn_prefix = clean(row[1]).upper()
        bundesland = clean(row[2])

        if not source_name and not zpn_prefix and not bundesland:
            continue

        if not source_name:
            raise ValueError(
                f"{path.name}, Zeile {row_number}: "
                "Netzbetreiber fehlt."
            )

        if not zpn_prefix:
            raise ValueError(
                f"{path.name}, Zeile {row_number}: "
                f"ZPN-Präfix fehlt für '{source_name}'."
            )

        if (
            len(zpn_prefix) != 8
            or not zpn_prefix.startswith("AT")
        ):
            raise ValueError(
                f"{path.name}, Zeile {row_number}: "
                f"Ungültiges ZPN-Präfix '{zpn_prefix}'."
            )

        if not bundesland:
            raise ValueError(
                f"{path.name}, Zeile {row_number}: "
                f"Bundesland fehlt für '{source_name}'."
            )

        operator, operator_created = get_or_create_operator(
            db,
            source_name,
        )

        if operator_created:
            created_operators += 1

        identifier = (
            db.query(NetworkOperatorIdentifier)
            .filter(
                NetworkOperatorIdentifier.energy_type
                == energy_type,
                NetworkOperatorIdentifier.zpn_prefix
                == zpn_prefix,
            )
            .first()
        )

        if identifier is None:
            identifier = NetworkOperatorIdentifier(
                network_operator_id=operator.id,
                energy_type=energy_type,
                zpn_prefix=zpn_prefix,
                bundesland=bundesland,
                active=True,
            )

            db.add(identifier)
            created_identifiers += 1

        else:
            changed = False

            if identifier.network_operator_id != operator.id:
                identifier.network_operator_id = operator.id
                changed = True

            if identifier.bundesland != bundesland:
                identifier.bundesland = bundesland
                changed = True

            if not identifier.active:
                identifier.active = True
                changed = True

            if changed:
                updated_identifiers += 1

        processed += 1

    workbook.close()

    return {
        "processed": processed,
        "created_operators": created_operators,
        "created_identifiers": created_identifiers,
        "updated_identifiers": updated_identifiers,
    }


def main() -> None:
    db = SessionLocal()

    try:
        total_processed = 0
        total_created_operators = 0
        total_created_identifiers = 0
        total_updated_identifiers = 0

        for config in FILES:
            path = find_file(config["pattern"])

            print(
                f"\nImportiere {config['energy_type']}: "
                f"{path.name}"
            )

            result = import_file(
                db,
                path=path,
                sheet_name=config["sheet"],
                energy_type=config["energy_type"],
            )

            total_processed += result["processed"]
            total_created_operators += result["created_operators"]
            total_created_identifiers += result["created_identifiers"]
            total_updated_identifiers += result["updated_identifiers"]

            print(
                f"  Verarbeitet: {result['processed']}\n"
                f"  Neue Netzbetreiber: "
                f"{result['created_operators']}\n"
                f"  Neue Kennungen: "
                f"{result['created_identifiers']}\n"
                f"  Aktualisierte Kennungen: "
                f"{result['updated_identifiers']}"
            )

        db.commit()

        print("\n===== IMPORT ABGESCHLOSSEN =====")
        print(f"Verarbeitet: {total_processed}")
        print(
            "Neue Netzbetreiber: "
            f"{total_created_operators}"
        )
        print(
            "Neue Kennungen: "
            f"{total_created_identifiers}"
        )
        print(
            "Aktualisierte Kennungen: "
            f"{total_updated_identifiers}"
        )

    except Exception:
        db.rollback()
        raise

    finally:
        db.close()


if __name__ == "__main__":
    main()
