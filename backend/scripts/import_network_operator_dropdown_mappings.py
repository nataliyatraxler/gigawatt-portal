from pathlib import Path

from openpyxl import load_workbook

from app.database.session import SessionLocal
from app.models.network_operator_identifier import NetworkOperatorIdentifier
from app.models.network_operator_dropdown_mapping import (
    NetworkOperatorDropdownMapping,
)


FILE_PATH = (
    Path.home()
    / "Downloads"
    / "STROM_Netzbetreiber_Programmierlogik_2026_HAUPTNETZBETREIBER.xlsx"
)

SHEET_NAME = "Dropdown_Mapping"


def normalize(value):
    if value is None:
        return ""
    return str(value).strip()


def main():
    if not FILE_PATH.exists():
        raise FileNotFoundError(f"Datei nicht gefunden: {FILE_PATH}")

    workbook = load_workbook(
        FILE_PATH,
        read_only=True,
        data_only=True,
    )

    if SHEET_NAME not in workbook.sheetnames:
        raise RuntimeError(
            f"Sheet '{SHEET_NAME}' wurde nicht gefunden."
        )

    worksheet = workbook[SHEET_NAME]

    rows = worksheet.iter_rows(values_only=True)

    headers = next(rows)

    header_map = {
        normalize(name): index
        for index, name in enumerate(headers)
        if name is not None
    }

    required_columns = [
        "Bundesland_Filter",
        "Netzbetreiber",
        "EC_Nummer_ZPN_Praefix",
        "Priorität",
        "Standard_sichtbar",
        "Quelle/Status",
    ]

    missing_columns = [
        column
        for column in required_columns
        if column not in header_map
    ]

    if missing_columns:
        raise RuntimeError(
            f"Fehlende Spalten: {missing_columns}"
        )

    db = SessionLocal()

    created = 0
    updated = 0
    unchanged = 0
    processed = 0
    missing_identifiers = []

    try:
        for row_number, row in enumerate(rows, start=2):
            bundesland = normalize(
                row[header_map["Bundesland_Filter"]]
            )

            operator_name = normalize(
                row[header_map["Netzbetreiber"]]
            )

            zpn_prefix = normalize(
                row[header_map["EC_Nummer_ZPN_Praefix"]]
            ).upper()

            priority_raw = row[header_map["Priorität"]]

            standard_visible_raw = normalize(
                row[header_map["Standard_sichtbar"]]
            ).upper()

            source_status = normalize(
                row[header_map["Quelle/Status"]]
            )

            if not bundesland and not zpn_prefix:
                continue

            if not bundesland or not zpn_prefix:
                raise RuntimeError(
                    f"Unvollständige Zeile {row_number}: "
                    f"Bundesland={bundesland!r}, "
                    f"EC-Präfix={zpn_prefix!r}"
                )

            try:
                priority = int(priority_raw)
            except (TypeError, ValueError):
                raise RuntimeError(
                    f"Ungültige Priorität in Zeile "
                    f"{row_number}: {priority_raw!r}"
                )

            standard_visible = standard_visible_raw == "JA"

            identifier = (
                db.query(NetworkOperatorIdentifier)
                .filter(
                    NetworkOperatorIdentifier.zpn_prefix
                    == zpn_prefix,
                    NetworkOperatorIdentifier.energy_type
                    == "Strom",
                )
                .first()
            )

            if identifier is None:
                missing_identifiers.append(
                    (
                        row_number,
                        bundesland,
                        operator_name,
                        zpn_prefix,
                    )
                )
                continue

            existing = (
                db.query(NetworkOperatorDropdownMapping)
                .filter(
                    NetworkOperatorDropdownMapping
                    .network_operator_identifier_id
                    == identifier.id,
                    NetworkOperatorDropdownMapping.bundesland
                    == bundesland,
                )
                .first()
            )

            processed += 1

            if existing is None:
                db.add(
                    NetworkOperatorDropdownMapping(
                        network_operator_identifier_id=identifier.id,
                        bundesland=bundesland,
                        priority=priority,
                        standard_visible=standard_visible,
                        source_status=source_status or None,
                    )
                )
                created += 1
                continue

            changed = False

            if existing.priority != priority:
                existing.priority = priority
                changed = True

            if existing.standard_visible != standard_visible:
                existing.standard_visible = standard_visible
                changed = True

            new_source_status = source_status or None

            if existing.source_status != new_source_status:
                existing.source_status = new_source_status
                changed = True

            if changed:
                updated += 1
            else:
                unchanged += 1

        if missing_identifiers:
            print("\n===== FEHLENDE IDENTIFIER =====")

            for item in missing_identifiers:
                print(
                    f"Zeile {item[0]} | "
                    f"{item[1]} | "
                    f"{item[2]} | "
                    f"{item[3]}"
                )

            raise RuntimeError(
                f"{len(missing_identifiers)} "
                "EC-Präfix(e) fehlen in der Datenbank. "
                "Import wurde abgebrochen."
            )

        db.commit()

        print("\n===== IMPORT ERFOLGREICH =====")
        print("Processed:", processed)
        print("Created:", created)
        print("Updated:", updated)
        print("Unchanged:", unchanged)

    except Exception:
        db.rollback()
        raise

    finally:
        db.close()
        workbook.close()


if __name__ == "__main__":
    main()
