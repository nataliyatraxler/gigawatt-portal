from decimal import Decimal
import re

from openpyxl import load_workbook

from app.database.session import SessionLocal
from app.models.sne_tariff import SneTariff


FILE_PATH = "data/Entgelte SNE-V.xlsx"
SHEET_NAME = "Novelle 2026"
YEAR = 2026


def to_decimal(value):
    if value is None or value == "" or value == "-":
        return None

    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))

    try:
        return Decimal(str(value).replace(",", "."))
    except Exception:
        return None


def parse_tariff_type(title):
    title = title.lower()

    if "nicht gem" in title:
        return "nicht_gemessen"

    if "unterbrechbar" in title:
        return "unterbrechbar"

    if "gem." in title or "gemessen" in title:
        return "gemessen"

    return "standard"


def main():
    workbook = load_workbook(
        FILE_PATH,
        data_only=True,
        read_only=True,
    )

    worksheet = workbook[SHEET_NAME]

    db = SessionLocal()

    current_level = None
    tariff_type = None
    imported = 0
    updated = 0

    try:
        for row in worksheet.iter_rows(values_only=True):
            level_title = row[2] if len(row) > 2 else None

            if isinstance(level_title, str) and level_title.startswith("Netzebene"):
                match = re.search(r"Netzebene\s+(\d+)", level_title)

                if match:
                    current_level = int(match.group(1))
                    tariff_type = parse_tariff_type(level_title)

                continue

            if current_level is None:
                continue

            network_area = row[3] if len(row) > 3 else None

            if not isinstance(network_area, str):
                continue

            network_area = network_area.strip()

            if not network_area:
                continue

            lp = to_decimal(row[4] if len(row) > 4 else None)
            ap = to_decimal(row[5] if len(row) > 5 else None)
            snap = to_decimal(row[6] if len(row) > 6 else None)

            dtap = None
            dnap = None
            network_loss = None

            if current_level == 7 and tariff_type == "nicht_gemessen":
                dtap = to_decimal(row[7] if len(row) > 7 else None)
                dnap = to_decimal(row[8] if len(row) > 8 else None)
                network_loss = to_decimal(row[9] if len(row) > 9 else None)
            else:
                network_loss = to_decimal(row[8] if len(row) > 8 else None)

                if network_loss is None:
                    network_loss = to_decimal(row[9] if len(row) > 9 else None)

            existing = (
                db.query(SneTariff)
                .filter(
                    SneTariff.year == YEAR,
                    SneTariff.network_level == current_level,
                    SneTariff.tariff_type == tariff_type,
                    SneTariff.network_area == network_area,
                )
                .first()
            )

            values = {
                "lp_cent": lp,
                "ap_cent_kwh": ap,
                "snap_cent_kwh": snap,
                "dtap_cent_kwh": dtap,
                "dnap_cent_kwh": dnap,
                "network_loss_cent_kwh": network_loss,
            }

            if existing:
                for key, value in values.items():
                    setattr(existing, key, value)

                updated += 1
            else:
                db.add(
                    SneTariff(
                        year=YEAR,
                        network_level=current_level,
                        tariff_type=tariff_type,
                        network_area=network_area,
                        **values,
                    )
                )

                imported += 1

        db.commit()

        print(f"Imported: {imported}")
        print(f"Updated: {updated}")

    except Exception:
        db.rollback()
        raise

    finally:
        db.close()


if __name__ == "__main__":
    main()
